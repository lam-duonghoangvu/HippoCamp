#!/usr/bin/env python3
"""
HippoCamp terminal API: list_files / return_txt / return_img / return_ori /
return_metadata / set_flags / hhelp.

Rebuilt from the public contract in docs/docker_api.md and the return examples
in agent/prompt_modules/config.py (the original image source is not public).
The same module backs both the rebuilt Docker image (run as `hippocamp_api`
via sudo) and local_bench/ (Docker-free), so both return identical JSON.

Paths are configured through env vars; the defaults are the container layout:
    HIPPOCAMP_DATA_ROOT      raw benchmark files        (/hippocamp/data)
    HIPPOCAMP_GOLD_ROOT      parsed-text JSON release   (/hippocamp/gold)
    HIPPOCAMP_METADATA_XLSX  metadata spreadsheet       (first *.xlsx in /hippocamp/metadata)
    HIPPOCAMP_OUTPUT_ROOT    rendered/copied outputs    (/hippocamp/output)
    HIPPOCAMP_STATE_DIR      feature-flag state         (/hippocamp/.private/state)

Paths shown to the caller always use the container form (/hippocamp/data/...,
/hippocamp/output/...), even in local mode, so agent-visible output matches.
"""

from __future__ import annotations

import base64
import email
import email.policy
import fnmatch
import functools
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional

CONTAINER_DATA = "/hippocamp/data"
CONTAINER_OUTPUT = "/hippocamp/output"

DATA_ROOT = Path(os.environ.get("HIPPOCAMP_DATA_ROOT", CONTAINER_DATA)).resolve()
GOLD_ROOT = Path(os.environ.get("HIPPOCAMP_GOLD_ROOT", "/hippocamp/gold")).resolve()
OUTPUT_ROOT = Path(os.environ.get("HIPPOCAMP_OUTPUT_ROOT", CONTAINER_OUTPUT))
STATE_DIR = Path(os.environ.get("HIPPOCAMP_STATE_DIR", "/hippocamp/.private/state"))
METADATA_DIR = Path("/hippocamp/metadata")

# Page render size. The only rendered sample in the released prompt config is
# a 2400x3200 RGB PNG, so text pages use that canvas and PDFs render 2400 px wide.
RENDER_WIDTH = 2400
TEXT_PAGE_SIZE = (2400, 3200)
TEXT_MARGIN = 120
TEXT_FONT_SIZE = 40

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
TEXT_EXTS = {".txt", ".md", ".csv", ".json", ".log", ".py", ".ics", ".ipynb", ".eml"}
OFFICE_EXTS = {".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls", ".odt", ".odp", ".ods", ".rtf"}
AV_EXTS = {".mp3", ".mp4", ".mkv", ".wav", ".m4a", ".mov", ".avi"}

HELP_TEXT = (
    "AVAILABLE COMMANDS:\n"
    "  return_txt <file_path>\n"
    "  return_img <file_path> [output_path] [--page N]\n"
    "  return_ori <file_path> [output_path]\n"
    "  return_metadata <file_path>\n"
    "  list_files [pattern]\n"
)


# ---------------------------------------------------------------------------
# helpers


def _ok(**fields) -> Dict[str, Any]:
    out: Dict[str, Any] = {"success": True}
    out.update(fields)
    out["error"] = None
    return out


def _err(message: str) -> Dict[str, Any]:
    return {"success": False, "error": message}


def _resolve(file_path: str) -> Optional[Path]:
    """Resolve a data-relative (or /hippocamp/data-absolute) path, rejecting escapes."""
    if not file_path:
        return None
    rel = file_path
    if rel.startswith(CONTAINER_DATA + "/"):
        rel = rel[len(CONTAINER_DATA) + 1:]
    candidate = (DATA_ROOT / rel).resolve()
    try:
        candidate.relative_to(DATA_ROOT)
    except ValueError:
        return None
    return candidate


def _rel(path: Path) -> str:
    return path.relative_to(DATA_ROOT).as_posix()


def _resolve_output(output_path: str, default_name: str) -> Optional[Path]:
    """Map an output path (container form or relative) into OUTPUT_ROOT; reject escapes."""
    root = OUTPUT_ROOT.resolve()
    if not output_path:
        rel = default_name
    elif output_path.startswith(CONTAINER_OUTPUT + "/"):
        rel = output_path[len(CONTAINER_OUTPUT) + 1:]
    elif os.path.isabs(output_path):
        return None
    else:
        rel = output_path
    candidate = (root / rel).resolve()
    try:
        candidate.relative_to(root)
    except ValueError:
        return None
    return candidate


def _container_output_path(path: Path) -> str:
    return f"{CONTAINER_OUTPUT}/{path.relative_to(OUTPUT_ROOT.resolve()).as_posix()}"


def _write_atomic(dest: Path, payload: bytes) -> None:
    # Write via temp file + rename so a symlink planted at `dest` is replaced,
    # never followed.
    dest.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=str(dest.parent), prefix=".tmp_")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(payload)
        os.chmod(tmp, 0o644)
        os.replace(tmp, dest)
    except Exception:
        if os.path.exists(tmp):
            os.unlink(tmp)
        raise


# ---------------------------------------------------------------------------
# feature flags


def _flags_file() -> Path:
    return STATE_DIR / "feature_flags.json"


def get_flags() -> Dict[str, int]:
    try:
        data = json.loads(_flags_file().read_text(encoding="utf-8"))
        return {"return_txt": int(data.get("return_txt", 1)), "return_img": int(data.get("return_img", 1))}
    except Exception:
        return {"return_txt": 1, "return_img": 1}


def set_flags(return_txt_flag: str = "1", return_img_flag: str = "1") -> Dict[str, Any]:
    if return_txt_flag not in ("0", "1") or return_img_flag not in ("0", "1"):
        return _err("usage: set_flags <return_txt 0|1> <return_img 0|1>")
    flags = {"return_txt": int(return_txt_flag), "return_img": int(return_img_flag)}
    try:
        STATE_DIR.mkdir(parents=True, exist_ok=True)
        _write_atomic(_flags_file(), json.dumps(flags).encode("utf-8"))
    except Exception as e:
        return _err(f"failed to set flags: {e}")
    return _ok(data=flags)


# ---------------------------------------------------------------------------
# list_files / return_txt / return_metadata / return_ori


def list_files(pattern: str = "") -> Dict[str, Any]:
    if not DATA_ROOT.is_dir():
        return _err(f"data root not found: {CONTAINER_DATA}")
    paths = sorted(_rel(p) for p in DATA_ROOT.rglob("*") if p.is_file())
    if pattern:
        if any(ch in pattern for ch in "*?["):
            paths = [p for p in paths if fnmatch.fnmatch(p, pattern) or fnmatch.fnmatch(os.path.basename(p), pattern)]
        else:
            # The WebUI route documents `?pattern=<glob-or-prefix>`.
            paths = [p for p in paths if p.startswith(pattern) or os.path.basename(p).startswith(pattern)]
    return _ok(count=len(paths), data=paths)


def _gold_path_for(rel_path: str) -> Path:
    return GOLD_ROOT / Path(rel_path).with_suffix(".json")


def _load_gold(rel_path: str) -> Optional[Dict[str, Any]]:
    gold_path = _gold_path_for(rel_path)
    if not gold_path.is_file():
        return None
    with gold_path.open("r", encoding="utf-8") as f:
        return json.load(f)


def return_txt(file_path: str) -> Dict[str, Any]:
    if not get_flags()["return_txt"]:
        return _err("return_txt is disabled by feature flags")
    resolved = _resolve(file_path)
    if resolved is None:
        return _err(f"invalid path: {file_path}")
    if not resolved.is_file():
        return _err(f"file not found: {file_path}")
    try:
        data = _load_gold(_rel(resolved))
    except Exception as e:
        return _err(f"failed to read parsed text: {e}")
    if data is None:
        return _err(f"no parsed text available for: {file_path}")
    return _ok(data=data)


def _metadata_xlsx() -> Optional[Path]:
    env = os.environ.get("HIPPOCAMP_METADATA_XLSX")
    if env:
        return Path(env)
    if METADATA_DIR.is_dir():
        found = sorted(METADATA_DIR.glob("*.xlsx"))
        if found:
            return found[0]
    return None


def _cell_str(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


@functools.lru_cache(maxsize=1)
def _load_metadata_rows() -> Dict[str, Dict[str, Any]]:
    xlsx = _metadata_xlsx()
    if not xlsx or not xlsx.is_file():
        return {}
    import openpyxl

    # data_only: some sheets store ID as `=ROW()-1`; use the cached value.
    wb = openpyxl.load_workbook(str(xlsx), read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        return {}
    idx = {str(h).strip(): i for i, h in enumerate(rows[0]) if h is not None}

    def col(row, *names):
        for name in names:
            i = idx.get(name)
            if i is not None and i < len(row):
                return row[i]
        return None

    out: Dict[str, Dict[str, Any]] = {}
    for n, row in enumerate(rows[1:], start=1):
        path = col(row, "FilePath", "file_path")
        if not path:
            continue
        row_id = col(row, "ID", "id")
        if row_id is None or (isinstance(row_id, str) and row_id.startswith("=")):
            row_id = n
        out[str(path)] = {
            "id": row_id,
            "file_path": str(path),
            "file_type": col(row, "FileType", "file_type"),
            "file_modality": col(row, "FileModality", "file_modality"),
            "creation_date": _cell_str(col(row, "creation_date")),
            "modification_date": _cell_str(col(row, "modification_date")),
            "latitude": col(row, "latitude"),
            "longitude": col(row, "longitude"),
            "location": col(row, "location"),
        }
    return out


def return_metadata(file_path: str) -> Dict[str, Any]:
    resolved = _resolve(file_path)
    if resolved is None:
        return _err(f"invalid path: {file_path}")
    if not resolved.is_file():
        return _err(f"file not found: {file_path}")
    rel = _rel(resolved)

    row = _load_metadata_rows().get(rel)
    if row:
        return _ok(metadata=row)

    # Fall back to the gold file_info when the spreadsheet lacks this path.
    try:
        info = (_load_gold(rel) or {}).get("file_info") or {}
    except Exception:
        info = {}
    if info:
        return _ok(metadata={
            "id": info.get("id"),
            "file_path": rel,
            "file_type": info.get("file_type"),
            "file_modality": info.get("file_modality"),
            "creation_date": info.get("creation_date"),
            "modification_date": info.get("modification_date"),
            "latitude": info.get("latitude"),
            "longitude": info.get("longitude"),
            "location": info.get("location"),
        })
    return _err(f"no metadata available for: {file_path}")


def return_ori(file_path: str, output_path: str = "") -> Dict[str, Any]:
    resolved = _resolve(file_path)
    if resolved is None:
        return _err(f"invalid path: {file_path}")
    if not resolved.is_file():
        return _err(f"file not found: {file_path}")
    try:
        payload = resolved.read_bytes()
    except Exception as e:
        return _err(f"failed to read file: {e}")
    if output_path:
        dest = _resolve_output(output_path, resolved.name)
        if dest is None:
            return _err(f"output_path must be under {CONTAINER_OUTPUT}: {output_path}")
        try:
            _write_atomic(dest, payload)
        except Exception as e:
            return _err(f"failed to write output: {e}")
    return _ok(
        file_path=f"{CONTAINER_DATA}/{_rel(resolved)}",
        file_b64=base64.b64encode(payload).decode("ascii"),
    )


# ---------------------------------------------------------------------------
# return_img


def _cache_dir() -> Path:
    d = Path(tempfile.gettempdir()) / f"hippocamp_render_{os.getuid()}"
    d.mkdir(mode=0o700, parents=True, exist_ok=True)
    return d


def _office_to_pdf(src: Path) -> Path:
    st = src.stat()
    key = hashlib.sha1(f"{src}|{st.st_size}|{st.st_mtime_ns}".encode()).hexdigest()
    work = _cache_dir() / key
    pdf = work / (src.stem + ".pdf")
    if pdf.is_file():
        return pdf
    work.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) not available")
    profile = _cache_dir() / "lo_profile"
    subprocess.run(
        [soffice, f"-env:UserInstallation=file://{profile}", "--headless",
         "--convert-to", "pdf", "--outdir", str(work), str(src)],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180, check=False,
        env={**os.environ, "HOME": str(_cache_dir())},
    )
    if not pdf.is_file():
        raise RuntimeError("document conversion failed")
    return pdf


def _render_pdf_page(pdf_path: Path, page: Optional[int]):
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        page_count = len(pdf)
        page_num = page or 1
        if page_num < 1 or page_num > page_count:
            raise ValueError(f"page {page_num} out of range (1-{page_count})")
        pg = pdf[page_num - 1]
        scale = RENDER_WIDTH / pg.get_width()
        image = pg.render(scale=scale).to_pil().convert("RGB")
        return image, page_count, page_num
    finally:
        pdf.close()


@functools.lru_cache(maxsize=1)
def _text_font():
    from PIL import ImageFont

    candidates = [
        ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", 2),  # SC; covers Latin too
        ("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", 0),
    ]
    for path, index in candidates:
        if os.path.exists(path):
            return ImageFont.truetype(path, TEXT_FONT_SIZE, index=index)
    return ImageFont.load_default()


def _read_text_for_render(src: Path) -> str:
    raw = src.read_bytes()
    if src.suffix.lower() == ".eml":
        msg = email.message_from_bytes(raw, policy=email.policy.default)
        lines = [f"{h}: {msg[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if msg[h]]
        body = msg.get_body(preferencelist=("plain", "html"))
        text = body.get_content() if body is not None else ""
        return "\n".join(lines) + "\n\n" + text
    return raw.decode("utf-8", errors="replace")


def _paginate_text(text: str) -> List[List[str]]:
    font = _text_font()
    max_w = TEXT_PAGE_SIZE[0] - 2 * TEXT_MARGIN
    line_h = int(TEXT_FONT_SIZE * 1.35)
    per_page = (TEXT_PAGE_SIZE[1] - 2 * TEXT_MARGIN) // line_h
    lines: List[str] = []
    for para in text.expandtabs(4).splitlines() or [""]:
        # Greedy word wrap; tokens wider than a line (long URLs, CJK runs)
        # fall back to character breaks.
        cur = ""
        for token in re.findall(r"\s+|\S+", para):
            if font.getlength(cur + token) <= max_w:
                cur += token
                continue
            if cur.strip():
                lines.append(cur.rstrip())
                cur = ""
            token = token if cur or not token.isspace() else ""
            for ch in token:
                if font.getlength(cur + ch) > max_w:
                    lines.append(cur)
                    cur = ch
                else:
                    cur += ch
        lines.append(cur.rstrip())
    return [lines[i:i + per_page] for i in range(0, max(len(lines), 1), per_page)] or [[]]


def _render_text_page(src: Path, page: Optional[int]):
    from PIL import Image, ImageDraw

    pages = _paginate_text(_read_text_for_render(src))
    page_num = page or 1
    if page_num < 1 or page_num > len(pages):
        raise ValueError(f"page {page_num} out of range (1-{len(pages)})")
    image = Image.new("RGB", TEXT_PAGE_SIZE, "white")
    draw = ImageDraw.Draw(image)
    font = _text_font()
    line_h = int(TEXT_FONT_SIZE * 1.35)
    y = TEXT_MARGIN
    for line in pages[page_num - 1]:
        draw.text((TEXT_MARGIN, y), line, fill="black", font=font)
        y += line_h
    return image, len(pages), page_num


def return_img(file_path: str, output_path: str = "", page: Optional[int] = None) -> Dict[str, Any]:
    if not get_flags()["return_img"]:
        return _err("return_img is disabled by feature flags")
    resolved = _resolve(file_path)
    if resolved is None:
        return _err(f"invalid path: {file_path}")
    if not resolved.is_file():
        return _err(f"file not found: {file_path}")

    suffix = resolved.suffix.lower()
    try:
        if suffix in IMAGE_EXTS:
            payload = resolved.read_bytes()
            page_count, page_num = 1, 1
            out_ext = suffix
        else:
            if suffix == ".pdf":
                image, page_count, page_num = _render_pdf_page(resolved, page)
            elif suffix in OFFICE_EXTS:
                image, page_count, page_num = _render_pdf_page(_office_to_pdf(resolved), page)
            elif suffix in TEXT_EXTS:
                image, page_count, page_num = _render_text_page(resolved, page)
            elif suffix in AV_EXTS:
                return _err(f"return_img not supported for audio/video file: {suffix}")
            else:
                return _err(f"return_img not supported for file type: {suffix or '(none)'}")
            buf = io.BytesIO()
            image.save(buf, format="PNG")
            payload = buf.getvalue()
            out_ext = ".png"
    except ValueError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"failed to render {file_path}: {e}")

    default_name = resolved.stem + (f"_page{page_num}" if page_count > 1 else "") + out_ext
    dest = _resolve_output(output_path, default_name)
    if dest is None:
        return _err(f"output_path must be under {CONTAINER_OUTPUT}: {output_path}")
    try:
        _write_atomic(dest, payload)
    except Exception as e:
        return _err(f"failed to write output: {e}")

    shown = _container_output_path(dest)
    b64 = base64.b64encode(payload).decode("ascii")
    return _ok(
        image_path=shown,
        image_paths=[shown],
        image_b64=b64,
        image_b64_list=[b64],
        page_count=page_count,
    )


def hhelp() -> Dict[str, Any]:
    return _ok(data=HELP_TEXT)


# ---------------------------------------------------------------------------
# CLI


def dispatch(argv: List[str]) -> Dict[str, Any]:
    if not argv:
        return _err("no command given")
    cmd, rest = argv[0], argv[1:]
    if cmd == "list_files":
        if len(rest) > 1:
            return _err("list_files accepts at most one pattern argument")
        return list_files(rest[0] if rest else "")
    if cmd in ("return_txt", "return_metadata"):
        if len(rest) != 1:
            return _err(f"usage: {cmd} <file_path>")
        return return_txt(rest[0]) if cmd == "return_txt" else return_metadata(rest[0])
    if cmd == "return_ori":
        if not rest or len(rest) > 2:
            return _err("usage: return_ori <file_path> [output_path]")
        return return_ori(rest[0], rest[1] if len(rest) > 1 else "")
    if cmd == "return_img":
        if not rest:
            return _err("usage: return_img <file_path> [output_path] [--page N]")
        file_path, output_path, page = rest[0], "", None
        i = 1
        while i < len(rest):
            if rest[i] == "--page":
                if i + 1 >= len(rest):
                    return _err("--page requires a number")
                try:
                    page = int(rest[i + 1])
                except ValueError:
                    return _err(f"invalid page number: {rest[i + 1]}")
                i += 2
                continue
            if output_path:
                return _err("usage: return_img <file_path> [output_path] [--page N]")
            output_path = rest[i]
            i += 1
        return return_img(file_path, output_path, page)
    if cmd == "set_flags":
        if len(rest) != 2:
            return _err("usage: set_flags <return_txt 0|1> <return_img 0|1>")
        return set_flags(rest[0], rest[1])
    if cmd == "hhelp":
        return hhelp()
    return _err(f"unknown command: {cmd}")


def render(result: Dict[str, Any]) -> str:
    return json.dumps(result, ensure_ascii=False, indent=2, default=str)


def main(argv: Optional[List[str]] = None) -> int:
    result = dispatch(sys.argv[1:] if argv is None else argv)
    print(render(result))
    return 0 if result.get("success") else 1


if __name__ == "__main__":
    sys.exit(main())
